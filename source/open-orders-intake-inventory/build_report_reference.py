#!/usr/bin/env python3
"""
Reference implementation (Python/openpyxl) of the "Open Orders - Order Intake - Inventory Analysis By MSA"
report, version 2.8. It mirrors build-open-orders-report.bas rule for rule (same rows, order, hidden rows,
formulas, chart and Run Log lines) so the VBA macro can be validated against it, and it is the fallback for
anyone who cannot run macros. Its add_open_orders_chart() is also used by run_macro_libreoffice.py to draw the
chart the macro records there.

Usage:
    python3 build_report_reference.py <input_folder> <output_xlsx> --size-class-file PATH
                                      [--msa-column NAME] [--inventory-month YYYY-MM-DD]
                                      [--show-cancelled] [--show-ssl]

Inputs (one folder with the Snowflake exports for all cities; Section 5.1 of the v2 plan):
    "Order Intake*.xlsx"        (required)  -> tab Order Intake
    "Inventory by Month*.xlsx"  (optional)  -> tab Inventory by Month (units by MSA, then the detail)
    "Open Order*.xlsx"          (optional)  -> read and logged only ("Open Orders" is accepted too)
    Every file of a kind is read and merged. Files starting with 001, containing "Analysis By MSA" or
    starting with ~$ are ignored. The MSA of a row comes from the MSA column (--msa-column, else the first header
    equal to MSA, containing MSA, equal to CITY, or containing CITY, METRO or MARKET), else from the file name.
    (In the code a row's MSA is called its city; the report calls it MSA.)
    --size-class-file is the master list (columns HP, Rated Operating Capacity, Name of machine, Size Class,
    matched by name with aliases). The macro also reads .xls; openpyxl cannot.

Rules (identical in the macro):
    * Order Intake keeps PRODUCT_TYPE_ABBR CTL, MX and SSL (written trimmed and upper-cased); other types are
      dropped and counted. Nothing is dropped for its status. CANCELLED rows and SSL rows are written and hidden
      by AutoFilter criteria unless --show-cancelled / --show-ssl.
    * Size class from MODEL_GROUP via the size class file (exact name, then the name's first word), the only
      source of a class. The allowed classes (v2.8) are the three tables of its Set-up tab (CTL size class, MX size
      class, SSL size class), in that order, top to bottom; without a Set-up tab, the 17 built in (CTL by rated
      operating capacity, MX by nominal tonnes, SSL by frame size). A model not in the file is UNMAPPED (bold, peach).
    * Sorted by size class, order status (ORDERED, Not yet processed, INVOICED, Completely processed, other,
      CANCELLED), ORDER_DATE newest first (date part, undated last), MSA, then export order.
    * Visible total (v2.6): the frozen title row of Order Intake holds "Total sum of ordered Quantity of visible
      selection:" (F1) and =SUBTOTAL(109,...) over the whole SUM(F.ORDERED_QUANTITY) column (G1), so the figure
      follows whatever filters the reader sets; the tab's note moves from D1 to H1.
    * Inventory keeps Compact Track Loaders, Mini Excavators, Skid-Steer Loaders; the month shown is
      --inventory-month (any day in it) or the latest month in the files. A "Units by MSA" table of SUMIFS
      formulas sits above the detail, which is filtered to that month (and to CTL and MX unless --show-ssl).
    * Dates stored as text are converted when they read as a date; anything else is left as it is and its rows
      go to the Run Log. Blank rows are ignored: rows read = written + dropped + blank.
    * Open Orders by MSA (v2.5): the ordered quantity of the Order Intake rows whose ORDER_STATUS is ORDERED or
      Not yet processed (as SUMIFS compares: whole text, ignoring case, not ignoring spaces), by MSA and product
      type, as live SUMIFS formulas over the Order Intake tab, and a clustered column chart of that table (one
      series per MSA; SSL on the chart only with --show-ssl). Plan: open-orders-by-msa-tab-implementation-plan.md.
"""
import argparse, os, re, sys, time, datetime as dt
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.filters import FilterColumn, CustomFilters, CustomFilter
from openpyxl.worksheet.datavalidation import DataValidation

VERSION = '2.8'
BUILDER = f'build_report_reference.py v{VERSION}'
# ----------------------------------------------------------------------------- the rules
# v2.7: the 17 size classes in report order (rank 1 to 17), three families; since v2.8 the built-in list, used when
# the size class file has no Set-up tab (SIZE_FAMILIES: its family of each)
SIZE_CLASSES = ['0<2000', '2000<2300', '2300<2600', '2600<3000', '3000<3500', '3500&Over',
                '>00-01', '>01-02', '>02-03', '>03-04', '>04-05', '>05-06', '>06-10',
                'EX Small', 'Small Frame', 'Med Frame', 'Large Frame']
UNMAPPED = 'UNMAPPED'
SIZE_FAMILIES = ['CTL'] * 6 + ['MX'] * 7 + ['SSL'] * 4
UNMAPPED_RANK, MAX_CLASSES = 99, 98               # UNMAPPED (anything else) ranks 99, so at most 98 classes
# v2.8: the Set-up tab of the size class file (sheet name, ignoring case, spaces, hyphens and underscores) and the
# headers of its three tables, read in report order CTL, MX, SSL
SETUP_SHEET_KEY = 'SETUP'
SETUP_HEADERS = [('CTL', 'CTL size class'), ('MX', 'MX size class'), ('SSL', 'SSL size class')]
CTL_FAMILY, MX_FAMILY, SSL_FAMILY = 'Track loaders (CTL)', 'Mini excavators (MX)', 'Skid-steer loaders (SSL)'
CTL_RULE_TEXT = 'by rated operating capacity in lb; the class is read from the size class file'
MX_RULE_TEXT = 'by nominal tonnes; the class is read from the size class file'
SSL_RULE_TEXT = 'by frame size; the class is read from the size class file'
STATUS_ORDER = ['ORDERED', 'Not yet processed', 'INVOICED', 'Completely processed', '(any other status)', 'CANCELLED']
STATUS_RANK = {'ORDERED': 1, 'NOT YET PROCESSED': 2, 'INVOICED': 3, 'COMPLETELY PROCESSED': 4, 'CANCELLED': 6}
OTHER_STATUS_RANK = 5
ABBRS = ['CTL', 'MX', 'SSL']
TYPE_NAMES = ['Compact Track Loaders', 'Mini Excavators', 'Skid-Steer Loaders']      # same order as ABBRS
OI_COLS = ['ORDER_DATE', 'REQUESTED_DATE', 'ORDER_STATUS', 'IS_OPEN_ORDER', 'SUM(F.ORDERED_QUANTITY)',
           'PRODUCT_TYPE_ABBR', 'MODEL_GROUP', 'MODEL_NUMBER', 'ORGANIZATION_NAME']
INV_COLS = ['PRODUCT_TYPE', 'MODEL_GROUP', 'MODEL_NUMBER', 'INVENTORY_DATE', 'TOTAL']
OI_HEADERS = ['Size Class', 'MSA'] + OI_COLS
SUM_HEADERS = ['MSA', 'CTL', 'MX', 'SSL', 'Total']
INV_EXTRA = ['ORGANIZATION_NAME', 'ORGANIZATION_CUSTOMER_NUMBER']    # optional; kept at the end of the detail (owner, 2026-10-12)
INV_HEADERS = ['MSA', 'PRODUCT_TYPE', 'MODEL_GROUP', 'MODEL_NUMBER', 'INVENTORY_DATE', 'Month', 'TOTAL'] + INV_EXTRA
KINDS = [('ORDER INTAKE', 'OI'), ('INVENTORY BY MONTH', 'INV'), ('OPEN ORDERS', 'OO'), ('OPEN ORDER', 'OO')]  # plural first
KIND_LABEL = {'OI': 'Order Intake', 'INV': 'Inventory by Month', 'OO': 'Open Order'}
CITY_TESTS = [lambda h: h == 'MSA', lambda h: 'MSA' in h, lambda h: h == 'CITY', lambda h: 'CITY' in h,
              lambda h: 'METRO' in h, lambda h: 'MARKET' in h]       # on headers without spaces and underscores
NO_CITY = '(no MSA)'
SC_ALIASES = [('hp', ['HP', 'HORSEPOWER']),             # v2.7: a rated operating capacity column is not read
              ('name', ['NAMEOFMACHINE', 'MACHINE', 'MODEL', 'MODELGROUP']),
              ('cls', ['SIZECLASS', 'CLASS'])]
MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September',
          'October', 'November', 'December']
# File Format layout: fixed rows (the macro uses the same)
FF_CLASS_ROW, FF_MONTH_COL = 3, 6
# v2.8: Table 1 has as many rows as classes, so the tables under it move: Table 2 two rows under the UNMAPPED row
# (FF_CLASS_ROW + n + 4; with the 17 built-in classes row 24 as in v2.7), Tables 3 and 4 each 9 rows further
FF_TABLE_GAP = 9
# Open Orders by MSA (v2.5): ORDER_STATUS as SUMIFS compares it (whole text, ignoring case, not ignoring spaces)
OO_SHEET = 'Open Orders by MSA'
OPEN_STATUSES = ['ORDERED', 'Not yet processed']
OO_WIDTHS = [36, 12, 12, 12, 12]
OO_CHART_TITLE = 'Open orders by product type and MSA (ordered quantity)'
# chart series fills in MSA order, wrapping after eight; the first is the report green
CHART_PALETTE = ['548235', '7F7F7F', '2F5597', 'C55A11', '7030A0', '70AD47', 'A5A5A5', '264478']
CHART_BORDER = 'D9D9D9'
MAX_ROWS_LISTED = 10                                  # the macro's MAX_ROWS_LISTED
# ----------------------------------------------------------------------------- style sheet (v2.4)
# Set like an engineering report, in green: Cambria (ships with Office), horizontal black rules only (no
# vertical lines), bold white headers on green with a medium rule above and a double rule under them, a medium
# rule between groups, visible body rows banded light green and white (conditional formatting, so the bands
# follow the filters), UNMAPPED Size Class cells peach, the month cell yellow. The Run Log is in Consolas.
FONT_NAME, LOG_FONT_NAME, BODY_SIZE = 'Cambria', 'Consolas', 10
BODY_FONT = Font(name=FONT_NAME, size=BODY_SIZE)
BOLD_FONT = Font(name=FONT_NAME, size=BODY_SIZE, bold=True)
HEAD_FONT = Font(name=FONT_NAME, size=BODY_SIZE, bold=True, color='FFFFFFFF')
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True)
NOTE_FONT = Font(name=FONT_NAME, size=BODY_SIZE, italic=True)
LOG_FONT = Font(name=LOG_FONT_NAME, size=BODY_SIZE)
SIDES = {k: Side(style=k, color='FF000000') for k in ('thin', 'medium', 'double')}
HEAD_FILL = PatternFill(fill_type='solid', fgColor='FF548235')      # header rows (white bold text)
BAND_FILL = PatternFill(fill_type='solid', start_color='FFE2EFDA', end_color='FFE2EFDA')   # every second visible row
FLAG_FILL = PatternFill(fill_type='solid', start_color='FFFCE4D6', end_color='FFFCE4D6')   # UNMAPPED Size Class cells
INPUT_FILL = PatternFill(fill_type='solid', fgColor='FFFFFFCC')     # the month cell
REPORT_TITLE = 'Open Orders - Order Intake - Inventory Analysis By MSA'   # the printed page header
HF_FONT, HF_SIZE = 'Cambria,Regular', 9
# widths as Excel's UI shows them (what the macro sets); openpyxl stores UI width + 0.71
OI_WIDTHS = [20.29, 18, 18.86, 22.57, 19.29, 21, 35.57, 22.29, 17.86, 24, 56.86]
INV_WIDTHS = [18, 25.14, 14, 16, 17, 16, 11.29, 56.86, 36]
FF_WIDTHS = [22, 26, 62, 32, 3, 20]
LOG_WIDTH = 150
# column alignment, one letter per column: L left, R right (dates, numbers), C centre (short values); L and R
# cells are indented one step, the cell padding of a printed table
OI_ALIGN, SUM_ALIGN, INV_ALIGN = 'LLRRLCRCLLL', 'LRRRR', 'LLLLRLRLR'
# Order Intake title row (v2.6): the visible total's label right-aligned in F1, its figure left-aligned in G1 (above
# SUM(F.ORDERED_QUANTITY)), so the two read as one line, and the tab's note in H1 instead of D1, so the label does
# not cut it off
OI_TOTAL_LABEL = 'Total sum of ordered Quantity of visible selection:'
OI_NOTE_COL = 8
OI_TOTAL_FONT = Font(name=FONT_NAME, size=12, bold=True)
HEAD_HEIGHT, ROW_HEIGHT = 20, 16
DATE_FMT, QTY_FMT, UNITS_FMT = 'm/d/yyyy', '0', '#,##0'
# text-date forms accepted, standing in for VBA CDate on a US-English PC (CDate accepts more); every form has a
# 4-digit year, as the macro requires
DATE_TEXT_FORMATS = ('%Y-%m-%d', '%Y-%m-%d %H:%M', '%Y-%m-%d %H:%M:%S', '%Y/%m/%d', '%Y/%m/%d %H:%M', '%Y/%m/%d %H:%M:%S',
                     '%m/%d/%Y', '%m/%d/%Y %H:%M', '%m/%d/%Y %H:%M:%S', '%m/%d/%Y %I:%M %p', '%m/%d/%Y %I:%M:%S %p',
                     '%m-%d-%Y', '%d-%b-%Y', '%d %b %Y', '%d %B %Y', '%b %d %Y', '%b %d, %Y', '%B %d %Y', '%B %d, %Y',
                     '%b %d, %Y %I:%M %p', '%B %d, %Y %I:%M %p')
ERROR_VALUES = {'#NULL!', '#DIV/0!', '#VALUE!', '#REF!', '#NAME?', '#NUM!', '#N/A', '#GETTING_DATA', '#SPILL!', '#CALC!'}

# ----------------------------------------------------------------------------- VBA look-alikes
def vb_trim(s):
    """VBA Trim: strips spaces only (not tabs or non-breaking spaces)."""
    return s.strip(' ')

def safe_str(v):
    """The macro's SafeStr: '' for empty and error cells, otherwise the text of the value."""
    if v is None or (isinstance(v, str) and v in ERROR_VALUES):
        return ''
    if isinstance(v, bool):
        return 'True' if v else 'False'
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)

def is_blank(v):
    return vb_trim(safe_str(v)) == ''

def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)

def fmt_num(v):
    """A number as VBA's CStr writes it (2300, not 2300.0)."""
    return str(int(v)) if float(v).is_integer() else str(v)

def plural(n, one, many=None):
    return f'{n} {one if n == 1 else (many or one + "s")}'

def norm_city(s):
    """MSA text: trimmed, runs of spaces collapsed (the macro's NormCity)."""
    s = vb_trim(s)
    while '  ' in s:
        s = s.replace('  ', ' ')
    return s

def norm_header(h):
    """A header compared ignoring case, spaces and underscores."""
    return h.upper().replace(' ', '').replace('_', '')

def month_key(d):
    return d.year * 12 + d.month - 1

def month_text(key):
    return '' if key is None else f'{MONTHS[key % 12]} {key // 12}'

def text_to_date(s):
    """The macro's TextToDate: a 'T' separator, a trailing 'Z' and fractional seconds (Snowflake's
    '.000') are tolerated, then the text must contain a 4-digit year and read as a date (not only a
    time). Returns a datetime or None."""
    s = vb_trim(s)
    if s[-1:] in ('Z', 'z'):
        s = s[:-1]
    if len(s) >= 11 and s[10] in 'Tt':
        s = s[:10] + ' ' + s[11:]
    p = s.rfind('.')
    if p > 0 and ':' in s and p > s.rfind(':') and s[p + 1:].isdigit():
        s = s[:p]
    if not re.search(r'\d{4}', s):
        return None
    for fmt in DATE_TEXT_FORMATS:
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None

def has_day(s):
    """The macro's HasDay: a text date in an export names a day, not only a month: three runs of digits, or two
    and a month name (CDate alone would read "August 2026" as the 1st)."""
    runs = len(re.findall(r'[0-9]+', s))
    return runs >= 3 or (runs == 2 and re.search(r'[A-Za-z]', s) is not None)

def convert_date(v):
    """(value, status) for a cell of a date column: 'date', 'converted' (text that reads as a date, value
    replaced), 'blank', or 'bad' (anything else, value left as it is)."""
    if v is None or (isinstance(v, str) and vb_trim(v) == ''):
        return v, 'blank'
    if isinstance(v, dt.date):
        return v, 'date'
    if isinstance(v, str):
        d = text_to_date(v) if has_day(v) else None
        return (d, 'converted') if d is not None else (v, 'bad')
    return v, 'bad'

def date_line(label, col_name, converted, bad_rows):
    """One Run Log line for a date column that held text, or None."""
    if not converted and not bad_rows:
        return None
    line = f'{label} {col_name}: {converted} text value(s) converted to dates, {len(bad_rows)} not a date and left as is'
    if bad_rows:
        shown = ', '.join(str(r) for r in bad_rows[:10]) + (', ...' if len(bad_rows) > 10 else '')
        line += f" ({'row' if len(bad_rows) == 1 else 'rows'} {shown})"
    return line

# ----------------------------------------------------------------------------- reading
class Table:
    """The macro's ReadSource: first worksheet; the header row is the first row with at least three non-blank
    cells; headers trimmed and upper-cased; blank rows below it are skipped (rows, nums) but counted (n_read)."""
    def __init__(self, headers, rows, nums, n_read):
        self.headers, self.rows, self.nums, self.n_read = headers, rows, nums, n_read

    def index(self, name):
        return self.headers.index(name) if name in self.headers else -1

def read_table(path):
    if path.lower().endswith('.xls'):
        raise ValueError(f'{os.path.basename(path)}: openpyxl cannot read .xls files; save it as .xlsx (the macro reads .xls)')
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    raw = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    wb.close()
    hdr = next((i for i, r in enumerate(raw) if sum(not is_blank(v) for v in r) >= 3), None)
    if hdr is None:
        return None
    headers = [vb_trim(safe_str(h)).upper() for h in raw[hdr]]
    rows, nums, last = [], [], hdr
    for i in range(hdr + 1, len(raw)):
        if any(not is_blank(v) for v in raw[i]):
            rows.append(list(raw[i]) + [None] * (len(headers) - len(raw[i])))
            nums.append(i + 1)             # worksheet row number, for the Run Log
            last = i
    return Table(headers, rows, nums, last - hdr)

# ----------------------------------------------------------------------------- the size class file
class SizeClassMaster:
    """The master list (plan 5.1.3): usable rows, warnings (one per row at most, in row order), problems; v2.8: the
    allowed classes, [(class, family)] in rank order, and the Run Log line that says where they came from."""
    def __init__(self, path):
        self.path, self.name = path, os.path.basename(path)
        self.entries, self.warnings, self.problems = [], [], []   # entries: (name key, first word, class)
        self.cols = {}
        self.classes, self.classes_line = [], ''

    def canon(self, raw):
        """The allowed spelling of a class (matched ignoring case after trimming), or None."""
        key = vb_trim(raw).upper()
        return next((c for c, _ in self.classes if c.upper() == key), None)

    def rank(self, cls):
        return next((i for i, (c, _) in enumerate(self.classes, start=1) if c == cls), UNMAPPED_RANK)

    @property
    def usable(self):
        return len(self.entries)

    def lookup(self, model_group):
        """Class for a data MODEL_GROUP: exact name first, then the name's first word; None if absent."""
        key = vb_trim(safe_str(model_group)).upper()
        for k, _, cls in self.entries:
            if k == key:
                return cls
        for _, w, cls in self.entries:
            if w == key:
                return cls
        return None

def read_setup(path):
    """The Set-up tab of a workbook as rows of values from A1 (openpyxl, cached values), or None if it has none."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = next((w for w in wb.worksheets if norm_header(w.title).replace('-', '') == SETUP_SHEET_KEY), None)
    rows = None if ws is None else [list(r) for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True)]
    wb.close()
    return rows

def load_classes(m, setup):
    """The allowed classes: the Set-up tab's three tables, each read top to bottom below its header (CTL, then MX, then
    SSL; blank cells skipped; a class seen before, ignoring case, or the word UNMAPPED warned and ignored), or the 17
    built in when there is no Set-up tab. False after a problem."""
    if setup is None:
        m.classes = list(zip(SIZE_CLASSES, SIZE_FAMILIES))
        m.classes_line = f'Size classes: the {len(m.classes)} built into the macro ({m.name} has no Set-up tab)'
        return True
    found = {}
    for r, row in enumerate(setup):
        for c, v in enumerate(row):
            key = norm_header(vb_trim(safe_str(v)))
            for fam, head in SETUP_HEADERS:
                if key == norm_header(head) and fam not in found:
                    found[fam] = (r, c)
    missing = [head for fam, head in SETUP_HEADERS if fam not in found]
    if missing:
        m.problems.append(f'The Set-up tab of {m.name} has no {", ".join(missing)} header (it needs one table of size '
                          f'classes each for CTL, MX and SSL).')
        return False
    first, counts = {}, {fam: 0 for fam, _ in SETUP_HEADERS}
    for fam, _ in SETUP_HEADERS:
        r0, c = found[fam]
        for r in range(r0 + 1, len(setup)):
            raw = vb_trim(safe_str(setup[r][c] if c < len(setup[r]) else None))
            if not raw:
                continue
            where = f'{openpyxl.utils.get_column_letter(c + 1)}{r + 1}'
            if raw.upper() == UNMAPPED:
                m.warnings.append(f'Set-up tab {where}: "{raw}" is not a size class name (UNMAPPED marks models not in '
                                  f'the file) - ignored')
            elif raw.upper() in first:
                m.warnings.append(f'Set-up tab {where}: "{raw}" is listed again (first in {first[raw.upper()]}) - ignored')
            elif len(m.classes) == MAX_CLASSES:
                m.problems.append(f'The Set-up tab of {m.name} lists more than {MAX_CLASSES} size classes.')
                return False
            else:
                first[raw.upper()] = where
                m.classes.append((raw, fam))
                counts[fam] += 1
    if not m.classes:
        m.problems.append(f'The Set-up tab of {m.name} lists no size classes (type them under CTL size class, MX size '
                          f'class and SSL size class).')
        return False
    m.classes_line = (f'Size classes: {len(m.classes)} from the Set-up tab ('
                      + ', '.join(f'{fam} {counts[fam]}' for fam, _ in SETUP_HEADERS) + ')')
    return True

def load_size_class_file(path):
    m = SizeClassMaster(path)
    if not os.path.isfile(path):
        m.problems.append(f'SizeClassFile not found: {path}. Click Choose size class file.')
        return m
    try:
        t = read_table(path)
        setup = read_setup(path)
    except Exception as exc:
        m.problems.append(f'The size class file could not be read: {exc}')
        return m
    if t is None:
        m.problems.append(f'The size class file {m.name} has no header row (a row with at least three filled cells).')
        return m
    for key, aliases in SC_ALIASES:
        m.cols[key] = next((i for i, h in enumerate(t.headers) if h and norm_header(h) in aliases), -1)
    missing = [label for key, label in (('name', 'Name of machine'), ('cls', 'Size Class')) if m.cols[key] < 0]
    if missing:
        m.problems.append(f'The size class file {m.name} has no {" or ".join(missing)} column '
                          f'(headers: {", ".join(h for h in t.headers if h)}).')
        return m
    if not load_classes(m, setup):
        return m
    seen = {}
    for row, num in zip(t.rows, t.nums):
        name = vb_trim(safe_str(row[m.cols['name']]))
        raw_cls = vb_trim(safe_str(row[m.cols['cls']]))
        where = f'Size class file row {num}' + (f' {name}' if name else '')
        if not name:
            m.warnings.append(f'{where}: no machine name - row skipped')
            continue
        if not raw_cls:
            m.warnings.append(f'{where}: no size class - row skipped')
            continue
        cls = m.canon(raw_cls)
        if cls is None:
            m.warnings.append(f'{where}: "{raw_cls}" is not a size class - row skipped')
            continue
        key = name.upper()
        if key in seen:
            m.warnings.append(f'{where}: listed again (first in row {seen[key]}) - row {seen[key]} is used')
            continue
        seen[key] = num
        m.entries.append((key, key.split(' ')[0], cls))
    if not m.entries:
        m.problems.append(f'The size class file {m.name} has no usable rows (a Name of machine and an allowed size '
                          f'class).')
    return m

# ----------------------------------------------------------------------------- inputs
class InputFile:
    def __init__(self, folder, name, kind, file_city):
        self.path, self.name, self.kind, self.file_city = os.path.join(folder, name), name, kind, file_city
        self.table, self.idx, self.missing = None, None, ''
        self.city_col, self.city_hdr, self.extra = -1, '', [-1, -1]
        self.written, self.dropped, self.blank_city = 0, {}, 0
        self.city_keys = set()
        self.conv, self.bad = [0, 0], [[], []]   # per date column (OI: ORDER_DATE, REQUESTED_DATE; INV: INVENTORY_DATE)

    @property
    def usable(self):
        return self.table is not None and not self.missing

    def city_text(self):
        if self.city_col >= 0:
            n = len(self.city_keys)
            return f'MSA column {self.city_hdr}, {n} {"MSA" if n == 1 else "MSAs"}'
        return f'MSA from file name ({self.file_city})'

    def summary(self):
        """The file's line (Run Log after 'File ', and the File Format build table)."""
        if self.table is None:
            return f'{self.name} is empty - not used'
        if self.missing:
            return f'{self.name} is missing column(s) {self.missing} - not used'
        if self.kind == 'OO':
            return f'{self.name}: {plural(len(self.table.rows), "row")}, {self.city_text()} - not used by this report'
        s = f'{self.name}: {plural(self.table.n_read, "row")} -> {self.written} written'
        nd = sum(self.dropped.values())
        if nd:
            s += f', {nd} dropped (' + ', '.join(f'{k} {v}' for k, v in self.dropped.items()) + ')'
        blank = self.table.n_read - len(self.table.rows)
        if blank:
            s += f', {blank} blank ignored'
        return s + '; ' + self.city_text()

def discover(folder):
    """The macro's DiscoverFiles: input files in alphabetical order (case-insensitive)."""
    names = sorted((f for f in os.listdir(folder)
                    if f.lower().endswith(('.xlsx', '.xlsm', '.xls')) and not f.startswith('~$')
                    and os.path.isfile(os.path.join(folder, f))), key=str.upper)
    files = []
    for f in names:
        base = vb_trim(os.path.splitext(f)[0])
        upper = base.upper()
        if upper.startswith('001') or 'ANALYSIS BY MSA' in upper:
            continue                              # an earlier report, not an input
        hit = next(((p, k) for p, k in KINDS if upper.startswith(p)), None)
        if hit is None:
            continue
        city = vb_trim(base[len(hit[0]):].split('_')[0]).lstrip('- ')
        files.append(InputFile(folder, f, hit[1], norm_city(city)))
    return files

def find_city_column(headers, city_column):
    if city_column:
        key = vb_trim(city_column).upper()
        return headers.index(key) if key in headers else -1
    for test in CITY_TESTS:
        for i, h in enumerate(headers):
            n = norm_header(h)
            if n and test(n):
                return i
    return -1

def read_inputs(files, city_column, problems):
    """Reads every file and resolves its city column. Problems as the macro's setup check words them."""
    required = {'OI': OI_COLS, 'INV': INV_COLS, 'OO': []}
    for f in files:
        try:
            f.table = read_table(f.path)
        except Exception as exc:
            problems.append(f'{f.name} could not be read: {exc}')
            continue
        if f.table is None:
            continue
        f.idx = [f.table.index(c) for c in required[f.kind]]
        f.missing = ', '.join(c for c, i in zip(required[f.kind], f.idx) if i < 0)
        if f.kind == 'INV':
            f.extra = [f.table.index(c) for c in INV_EXTRA]
        f.city_col = find_city_column(f.table.headers, city_column)
        if f.city_col >= 0:
            f.city_hdr = f.table.headers[f.city_col]
        elif city_column:
            problems.append(f'MSAColumn "{vb_trim(city_column)}" is not a header of {f.name} '
                            f'(headers: {", ".join(h for h in f.table.headers if h)}).')
        elif not f.file_city:
            problems.append(f'No MSA column found in {f.name} (headers: {", ".join(h for h in f.table.headers if h)}). '
                            f'Type the header name into MSAColumn.')

class Data:
    """Everything a build needs, prepared from the files (the macro's PrepareData)."""
    def __init__(self):
        self.registry = {}            # upper-case city -> first spelling seen
        self.oi, self.inv = [], []    # row records in export order (files alphabetical, rows in order)
        self.unmapped = {}            # city key -> [model groups] in export order

    def city(self, f, row):
        if f.city_col >= 0:
            c = norm_city(safe_str(row[f.city_col]))
            if not c:
                f.blank_city += 1
                c = NO_CITY
        else:
            c = f.file_city
        key = c.upper()
        f.city_keys.add(key)
        if f.kind != 'OO':
            self.registry.setdefault(key, c)
        return key

def prepare(files, master):
    d = Data()
    for f in (f for f in files if f.kind == 'OI' and f.usable):
        for row, num in zip(f.table.rows, f.table.nums):
            raw_abbr = vb_trim(safe_str(row[f.idx[5]]))
            abbr = raw_abbr.upper()
            if abbr not in ABBRS:
                label = raw_abbr or '(blank)'
                f.dropped[label] = f.dropped.get(label, 0) + 1
                continue
            f.written += 1
            city = d.city(f, row)
            vals = [row[i] for i in f.idx]
            vals[5] = abbr
            for k in (0, 1):                       # ORDER_DATE, REQUESTED_DATE
                vals[k], st = convert_date(vals[k])
                if st == 'converted':
                    f.conv[k] += 1
                elif st == 'bad':
                    f.bad[k].append(num)
            mg = vb_trim(safe_str(row[f.idx[6]]))
            cls = master.lookup(mg) if master is not None else None
            if cls is None:                        # v2.7: not in the file = UNMAPPED, whatever the type or name
                cls = UNMAPPED
                groups = d.unmapped.setdefault(city, [])
                if (mg or '(blank)') not in groups:
                    groups.append(mg or '(blank)')
            status = safe_str(vals[2])
            day = vals[0].toordinal() if isinstance(vals[0], dt.date) else None
            d.oi.append({'seq': len(d.oi), 'cls': cls, 'city': city, 'vals': vals, 'abbr': abbr, 'status': status,
                         'status_rank': STATUS_RANK.get(vb_trim(status).upper(), OTHER_STATUS_RANK), 'day': day})
    upper_types = [t.upper() for t in TYPE_NAMES]
    for f in (f for f in files if f.kind == 'INV' and f.usable):
        for row, num in zip(f.table.rows, f.table.nums):
            raw_type = vb_trim(safe_str(row[f.idx[0]]))
            if raw_type.upper() not in upper_types:
                label = raw_type or '(blank)'
                f.dropped[label] = f.dropped.get(label, 0) + 1
                continue
            f.written += 1
            ti = upper_types.index(raw_type.upper())
            city = d.city(f, row)
            v, st = convert_date(row[f.idx[3]])
            if st == 'converted':
                f.conv[0] += 1
            elif st == 'bad':
                f.bad[0].append(num)
            date = isinstance(v, dt.date)
            extra = [row[i] if i >= 0 else None for i in f.extra]
            d.inv.append({'city': city, 'type': ti, 'vals': [TYPE_NAMES[ti], row[f.idx[1]], row[f.idx[2]], v, row[f.idx[4]]] + extra,
                          'month': month_key(v) if date else None, 'day': v.toordinal() if date else None})
    for f in (f for f in files if f.kind == 'OO' and f.table is not None):
        for row in f.table.rows:
            d.city(f, row)
    return d

# ----------------------------------------------------------------------------- writing helpers
def put(ws, row, col, value):
    """Write a value; text starting with '=' stays text, as the macro's SafeCellText keeps it."""
    c = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith('='):
        c.data_type = 's'
    return c

def title(ws, text_a1, note, ncols, note_col=4):
    """Title in A1, the note in D1 (Order Intake: H1), a medium rule under row 1 across the sheet's tables."""
    ws['A1'] = text_a1
    ws['A1'].font = TITLE_FONT
    put(ws, 1, note_col, note).font = NOTE_FONT
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.border = Border(bottom=SIDES['medium'])

def caption(ws, row, col, label, text_):
    """A table caption: 'Table 1.' in bold, then the description."""
    ws.cell(row=row, column=col).value = CellRichText(
        TextBlock(InlineFont(rFont=FONT_NAME, sz=BODY_SIZE, b=True), label),
        TextBlock(InlineFont(rFont=FONT_NAME, sz=BODY_SIZE), ' ' + text_))
    ws.cell(row=row, column=col).font = BODY_FONT

def note(ws, row, col, text_):
    put(ws, row, col, text_).font = NOTE_FONT

def align(ws, first, last, spec, col=1):
    for r in range(first, last + 1):
        for i, a in enumerate(spec):
            ws.cell(row=r, column=col + i).alignment = Alignment(
                horizontal={'L': 'left', 'R': 'right', 'C': 'center'}[a], vertical='center', indent=0 if a == 'C' else 1)

def header_row(ws, row, col, names, spec):
    for i in range(len(spec)):
        c = ws.cell(row=row, column=col + i)
        if i < len(names):
            c.value = names[i]
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    align(ws, row, row, spec, col)
    ws.row_dimensions[row].height = HEAD_HEIGHT

def grid(ws, first, last, ncols, lines=None, col=1):
    """Horizontal black rules over rows first..last and ncols columns from col, no vertical lines. lines maps a row
    to the rule above it ('double' under a header, 'medium' between groups); every other rule is thin."""
    lines = lines or {}
    for r in range(first, last + 1):
        top, bottom = SIDES[lines.get(r, 'thin')], SIDES[lines.get(r + 1, 'thin')]
        for c in range(col, col + ncols):
            ws.cell(row=r, column=c).border = Border(top=top, bottom=bottom)

def band(ws, first, last, col, ncols):
    """Every second visible row of first..last light green: the count of visible rows down to this one
    (SUBTOTAL 103 skips hidden rows) is even. Conditional formatting, so the bands follow any filter."""
    if last < first:
        return
    a, z = openpyxl.utils.get_column_letter(col), openpyxl.utils.get_column_letter(col + ncols - 1)
    ws.conditional_formatting.add(f'{a}{first}:{z}{last}', FormulaRule(
        formula=[f'MOD(SUBTOTAL(103,${a}${first}:${a}{first}),2)=0'], fill=BAND_FILL, stopIfTrue=False))

def ruled_table(ws, head, last, spec, groups=(), col=1, unbanded=0):
    """A header row at head and its body to last: horizontal rules, a medium rule above the header and a double rule
    under it, medium rules above groups, the columns aligned by spec, body rows ROW_HEIGHT high and banded
    (all but the last unbanded rows)."""
    lines = {head: 'medium', head + 1: 'double'}
    lines.update({g: 'medium' for g in groups if g > head + 1})
    grid(ws, head, max(last, head), len(spec), lines, col)
    align(ws, head + 1, last, spec, col)
    for r in range(head + 1, last + 1):
        ws.row_dimensions[r].height = ROW_HEIGHT
    band(ws, head + 1, last - unbanded, col, len(spec))

def finish_fonts(ws, font=BODY_FONT):
    """Every written cell not yet styled gets the body font."""
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and cell.font.name != font.name:
                cell.font = font

def widths(ws, ui_widths):
    for i, w in enumerate(ui_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = round(w + 0.71, 2)

def print_setup(ws, built, title_rows=None, pages_tall=0):
    """Landscape, one page wide (pages_tall 0 = as many as needed); page header and footer as in the LaTeX
    documents: report name and build date at the top, tab name and page number at the bottom."""
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, pages_tall
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75, header=0.3, footer=0.3)
    for part, value in ((ws.oddHeader.left, REPORT_TITLE), (ws.oddHeader.right, f'Built {built:%Y-%m-%d}'),
                        (ws.oddFooter.left, '&A'), (ws.oddFooter.center, 'Page &P of &N')):
        part.text, part.font, part.size = value, HF_FONT, HF_SIZE
    if title_rows:
        ws.print_title_rows = title_rows

def city_names(cities):
    return f'{len(cities)}: ' + ', '.join(cities) if cities else 'none found'

def is_open_status(status):
    """ORDER_STATUS ORDERED or Not yet processed as SUMIFS compares it: whole text, ignoring case, not spaces."""
    return status.upper() in [s.upper() for s in OPEN_STATUSES]

class OpenOrders:
    """The macro's open orders tally (Run Log lines and the File Format build table), over the same rows the
    SUMIFS on the Open Orders by MSA tab count, all three product types whatever ShowSSL says."""
    def __init__(self, rows):
        self.rows, self.ordered, self.units = {}, {}, {}
        self.total_rows, self.total_units = 0, [0, 0, 0]
        self.spaced, self.spaced_texts, self.disagree = 0, [], 0
        for r in sorted(rows, key=lambda r: r['seq']):          # export order, as the macro reads mOI
            key, open_ = r['city'], is_open_status(r['status'])
            if open_:
                self.rows[key] = self.rows.get(key, 0) + 1
                self.ordered[key] = self.ordered.get(key, 0) + (r['status'].upper() == 'ORDERED')
                u = self.units.setdefault(key, [0, 0, 0])
                if is_number(r['vals'][4]):
                    u[ABBRS.index(r['abbr'])] += r['vals'][4]
                    self.total_units[ABBRS.index(r['abbr'])] += r['vals'][4]
                self.total_rows += 1
            elif vb_trim(r['status']).upper() in [s.upper() for s in OPEN_STATUSES]:
                self.spaced += 1
                if r['status'] not in self.spaced_texts:
                    self.spaced_texts.append(r['status'])
            flag = r['vals'][3]
            if isinstance(flag, bool) and flag != open_:
                self.disagree += 1

    def city_line(self, city, key):
        n = self.rows.get(key, 0)
        if not n:
            return f'{city}: no open orders (ORDERED or Not yet processed)'
        u, o = self.units[key], self.ordered[key]
        return (f'{city}: Open orders {plural(n, "row")} (ORDERED {o}, Not yet processed {n - o}): CTL {fmt_num(u[0])}, '
                f'MX {fmt_num(u[1])}, SSL {fmt_num(u[2])} units')

    def tail_lines(self):
        out = []
        if self.total_rows == 0:
            out.append('Open orders: no rows with ORDER_STATUS ORDERED or Not yet processed - the Open Orders by MSA tab shows zeros')
        if self.spaced:
            listed = ', '.join(f'"{t}"' for t in self.spaced_texts[:MAX_ROWS_LISTED])
            if len(self.spaced_texts) > MAX_ROWS_LISTED:
                listed += ', ...'
            out.append(f'WARNING: {plural(self.spaced, "row")} have ORDER_STATUS ORDERED or Not yet processed with extra spaces '
                       f'and are not counted on the Open Orders by MSA tab: {listed}')
        if self.disagree:
            out.append(f'Open orders: {plural(self.disagree, "row")} where IS_OPEN_ORDER disagrees with the status (TRUE with '
                       f'another status, or FALSE with ORDERED / Not yet processed)')
        return out

    def build_text(self):
        u = self.total_units
        return (f'ORDERED and Not yet processed: {plural(self.total_rows, "row")}, {fmt_num(sum(u))} units (CTL {fmt_num(u[0])}, '
                f'MX {fmt_num(u[1])}, SSL {fmt_num(u[2])}), on the Open Orders by MSA tab')

def chart_text(size, bold=False):
    """Cambria at size points for a chart element's text."""
    cp = CharacterProperties(latin=DrawingFont(typeface=FONT_NAME), sz=size * 100, b=bold)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])

def chart_title(text_, size, bold=False):
    cp = CharacterProperties(latin=DrawingFont(typeface=FONT_NAME), sz=size * 100, b=bold)
    return Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text_)])])),
                 overlay=False)

def add_open_orders_chart(ws, nc, nt, anchor_row):
    """The clustered column chart of the Open Orders by MSA table (the macro's DrawOpenOrdersChart): one series per
    MSA row (5 to 4 + nc) named from column A, categories CTL, MX (and SSL when nt = 3) from row 4, the Total column
    and row left out; 700 x 320 pt at A<anchor_row>. Called once, after the sheet is written: openpyxl 3.1 reads a
    chart back when it loads a workbook, so drawing on a reloaded file that already has one would make two."""
    ch = BarChart()
    ch.type, ch.grouping, ch.gapWidth = 'col', 'clustered', 80
    ch.add_data(Reference(ws, min_col=1, max_col=1 + nt, min_row=5, max_row=4 + nc), from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, max_col=1 + nt, min_row=4, max_row=4))
    ch.title = chart_title(OO_CHART_TITLE, 12, True)
    ch.legend.position = 'b'
    ch.legend.txPr = chart_text(10)
    ch.y_axis.title = chart_title('Ordered quantity', 10)
    ch.y_axis.number_format = '#,##0'
    ch.y_axis.majorGridlines = ChartLines()
    ch.x_axis.majorGridlines = None
    for axis in (ch.x_axis, ch.y_axis):
        axis.delete = False                         # openpyxl 3.1 hides axes unless told otherwise
        axis.txPr = chart_text(10)
    for i, s in enumerate(ch.series):
        s.graphicalProperties.solidFill = CHART_PALETTE[i % len(CHART_PALETTE)]
        s.graphicalProperties.line.solidFill = CHART_PALETTE[i % len(CHART_PALETTE)]
        s.dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False)
        s.dLbls.numFmt = '#,##0'
        s.dLbls.txPr = chart_text(10)
    ch.graphical_properties = GraphicalProperties(noFill=True, ln=LineProperties(solidFill=CHART_BORDER, w=9525))
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True)
    ch.width, ch.height = 24.7, 11.3                 # cm: 700 x 320 pt
    ws.add_chart(ch, f'A{anchor_row}')

# ----------------------------------------------------------------------------- the build
class Build:
    def __init__(self, folder, opts):
        self.folder, self.opts, self.problems, self.log = folder, opts, [], []
        self.started = time.time()
        self.master = load_size_class_file(opts.size_class_file)
        self.problems += self.master.problems
        self.files = discover(folder)
        if not any(f.kind == 'OI' for f in self.files):
            self.problems.append(f'No Order Intake files in {folder}. Name them like "Order Intake 2026.xlsx" '
                                 f'(optional: "Inventory by Month 2026.xlsx", "Open Orders 2026.xlsx").')
        read_inputs(self.files, opts.city_column, self.problems)
        if self.problems:
            return
        self.data = prepare(self.files, self.master)
        d = self.data
        self.cities = sorted(d.registry.values(), key=str.upper)
        self.city_rank = {c.upper(): i for i, c in enumerate(self.cities)}
        months = {r['month'] for r in d.inv if r['month'] is not None}
        if opts.inventory_month:
            self.month = month_key(opts.inventory_month)
        else:
            self.month = max(months) if months else None
        if self.month is not None:
            months.add(self.month)
        self.months = sorted(months)
        d.oi.sort(key=lambda r: (self.master.rank(r['cls']), r['status_rank'], r['day'] is None,
                                 -(r['day'] or 0), self.city_rank[r['city']]))
        d.inv.sort(key=lambda r: (self.city_rank[r['city']], r['type'], r['day'] is None, -(r['day'] or 0)))
        for r in d.oi:
            r['hidden'] = ((not opts.show_cancelled and r['status'].upper() == 'CANCELLED')
                           or (not opts.show_ssl and r['abbr'] == 'SSL'))
        for r in d.inv:
            r['hidden'] = ((self.month is not None and r['month'] != self.month)
                           or (not opts.show_ssl and r['type'] == 2))

    def kind_files(self, kind):
        return [f for f in self.files if f.kind == kind]

    def options_line(self):
        o = self.opts
        return (f'Options: InventoryMonth={month_text(month_key(o.inventory_month)) if o.inventory_month else "latest in file"}, '
                f'ShowCancelled={"TRUE" if o.show_cancelled else "FALSE"}, ShowSSL={"TRUE" if o.show_ssl else "FALSE"}, '
                f'MSAColumn={vb_trim(o.city_column) if o.city_column else "auto"}, SizeClassFile={self.master.name}')

    def run_log_lines(self):
        d, log = self.data, []
        for f in self.kind_files('OI') + self.kind_files('INV'):
            log.append('File ' + f.summary())
        log.append(f'Size class file {self.master.name}: {plural(self.master.usable, "usable row")}, '
                   f'{plural(len(self.master.warnings), "warning")}')
        log.append(self.master.classes_line)
        log += self.master.warnings
        mtext = month_text(self.month) or '(no month)'
        self.open = OpenOrders(d.oi)
        for city in self.cities:
            key = city.upper()
            rows = [r for r in d.oi if r['city'] == key]
            if rows:
                n = {a: sum(r['abbr'] == a for r in rows) for a in ABBRS}
                canc = sum(r['status_rank'] == 6 for r in rows)
                line = (f'{city}: Order Intake {plural(len(rows), "row")} (CTL {n["CTL"]}, MX {n["MX"]}, SSL {n["SSL"]}; '
                        f'{canc} CANCELLED), {sum(r["hidden"] for r in rows)} hidden by the filters')
                if d.unmapped.get(key):
                    line += '; UNMAPPED model groups: ' + ', '.join(d.unmapped[key])
                log.append(line)
            else:
                log.append(f'{city}: no Order Intake rows')
            inv = [r for r in d.inv if r['city'] == key]
            if inv:
                units = [sum(r['vals'][4] for r in inv if r['type'] == t and r['month'] == self.month
                             and self.month is not None and is_number(r['vals'][4])) for t in range(3)]
                log.append(f'{city}: Inventory {plural(len(inv), "row")}; {mtext}: CTL {fmt_num(units[0])}, '
                           f'MX {fmt_num(units[1])}, SSL {fmt_num(units[2])} units')
            else:
                log.append(f'{city}: no inventory rows')
            log.append(self.open.city_line(city, key))
        for f in self.kind_files('OO'):
            log.append('File ' + f.summary())
        for f in self.kind_files('OI') + self.kind_files('INV'):
            if not f.usable:
                continue
            cols = ('ORDER_DATE', 'REQUESTED_DATE') if f.kind == 'OI' else ('INVENTORY_DATE',)
            for k, col in enumerate(cols):
                line = date_line(f.name, col, f.conv[k], f.bad[k])
                if line:
                    log.append(line)
        for kind in ('OI', 'INV', 'OO'):
            names = [f.name for f in self.kind_files(kind)]
            if len(names) > 1:
                log.append(f'{len(names)} {KIND_LABEL[kind]} files found, all are read: {", ".join(names)}')
        for f in self.files:
            if f.blank_city and f.kind != 'OO':
                log.append(f'{f.name}: {plural(f.blank_city, "row")} with a blank MSA - shown as {NO_CITY}')
        if not any(f.usable for f in self.kind_files('INV')):
            log.append('Inventory: no Inventory by Month file - the summary shows zeros')
        elif self.month is None:
            log.append('Inventory: no valid INVENTORY_DATE values - the summary shows zeros')
        elif not any(r['month'] == self.month for r in d.inv):
            log.append(f'Inventory: no rows for {month_text(self.month)} - the summary shows zeros')
        log += self.open.tail_lines()
        return log

    # ------------------------------------------------------------------------- sheets
    def write(self, out_path):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        self.built = dt.datetime.now()
        self.log = self.run_log_lines()
        self.file_format(wb.create_sheet('File Format'))
        self.order_intake(wb.create_sheet('Order Intake'))
        self.inventory(wb.create_sheet('Inventory by Month'))
        self.open_orders(wb.create_sheet(OO_SHEET))
        self.run_log(wb.create_sheet('Run Log'))
        for ws, rows, tall in zip(wb.worksheets, (None, '1:2', f'{self.inv_head}:{self.inv_head}', None, None), (1, 0, 0, 1, 0)):
            ws.sheet_view.showGridLines = False
            print_setup(ws, self.built, rows, tall)
        wb.save(out_path)

    def file_format(self, ws):
        o = self.opts
        title(ws, 'File Format', 'Layout and rules of this report.', FF_MONTH_COL)
        r = FF_CLASS_ROW
        caption(ws, r, 1, 'Table 1.', 'Size classes, in report order.')
        header_row(ws, r + 1, 1, ['Class', 'Family', 'Rule'], 'LLL')
        family = {'CTL': (CTL_FAMILY, CTL_RULE_TEXT), 'MX': (MX_FAMILY, MX_RULE_TEXT), 'SSL': (SSL_FAMILY, SSL_RULE_TEXT)}
        classes = self.master.classes
        rows = [(c, *family[f]) for c, f in classes]
        rows += [(UNMAPPED, 'Any', 'model not in the size class file; add it there and build again')]
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                put(ws, r + 2 + i, 1 + j, v)
        ws.cell(row=r + 1 + len(rows), column=1).font = BOLD_FONT
        ws.cell(row=r + 1 + len(rows), column=1).fill = FLAG_FILL
        # the UNMAPPED row is left out of the banding, or on an even row the green would hide it (v2.7); a medium rule
        # above the first class of each family after the first
        starts = {r + 2 + i for i in range(1, len(classes)) if classes[i][1] != classes[i - 1][1]}
        ruled_table(ws, r + 1, r + 1 + len(rows), 'LLL', starts, unbanded=1)
        r = status_row = r + 1 + len(rows) + 2
        caption(ws, r, 1, 'Table 2.', 'Order status rank inside a size class.')
        header_row(ws, r + 1, 1, ['Rank', 'ORDER_STATUS'], 'CL')
        for i, st in enumerate(STATUS_ORDER):
            ws.cell(row=r + 2 + i, column=1, value=i + 1)
            put(ws, r + 2 + i, 2, st)
        ruled_table(ws, r + 1, r + 1 + len(STATUS_ORDER), 'CL')
        r = filter_row = status_row + FF_TABLE_GAP
        caption(ws, r, 1, 'Table 3.', 'Default filters.')
        header_row(ws, r + 1, 1, ['Tab', 'Column', 'Shown', 'Config setting'], 'LLLL')
        tf = lambda b: 'TRUE' if b else 'FALSE'
        mtext = month_text(self.month) or '(no inventory month)'
        filters = [
            ('Order Intake', 'ORDER_STATUS', 'all' if o.show_cancelled else 'all except CANCELLED', f'ShowCancelled = {tf(o.show_cancelled)}'),
            ('Order Intake', 'PRODUCT_TYPE_ABBR', 'all (CTL, MX, SSL)' if o.show_ssl else 'CTL, MX', f'ShowSSL = {tf(o.show_ssl)}'),
            ('Inventory by Month', 'Month', mtext,
             'InventoryMonth = ' + (month_text(month_key(o.inventory_month)) if o.inventory_month else 'latest in file')),
            ('Inventory by Month', 'PRODUCT_TYPE', ('all (' + ', '.join(TYPE_NAMES) + ')') if o.show_ssl else ', '.join(TYPE_NAMES[:2]),
             f'ShowSSL = {tf(o.show_ssl)}'),
            (OO_SHEET, 'PRODUCT_TYPE_ABBR', 'CTL, MX, SSL on the chart' if o.show_ssl else 'CTL, MX on the chart',
             f'ShowSSL = {tf(o.show_ssl)}')]
        for i, row in enumerate(filters):
            for j, v in enumerate(row):
                put(ws, r + 2 + i, 1 + j, v)
        ruled_table(ws, r + 1, r + 1 + len(filters), 'LLLL')
        note(ws, r + 2 + len(filters), 1, 'Note. Clear a filter to see the rows it hides.')
        r = filter_row + FF_TABLE_GAP
        caption(ws, r, 1, 'Table 4.', 'This build.')
        header_row(ws, r + 1, 1, ['Item', 'Value'], 'LLLL')
        items = [('Built', self.built.strftime('%Y-%m-%d %H:%M')), ('Builder', BUILDER), ('Input folder', os.path.abspath(self.folder))]
        items += [(KIND_LABEL[f.kind] + ' file', f.summary()) for kind in ('OI', 'INV', 'OO') for f in self.kind_files(kind)]
        items += [('Size class file', f'{os.path.abspath(self.master.path)} ({plural(self.master.usable, "usable row")})'),
                  ('Inventory month', month_text(self.month) or '(none)'), ('MSAs', city_names(self.cities)),
                  ('Open orders', self.open.build_text()), ('Options', self.options_line()[len('Options: '):])]
        for i, (k, v) in enumerate(items):
            put(ws, r + 2 + i, 1, k)
            put(ws, r + 2 + i, 2, v)
        ruled_table(ws, r + 1, r + 1 + len(items), 'LLLL')
        caption(ws, FF_CLASS_ROW, FF_MONTH_COL, 'Table 5.', 'Months in this build.')
        for i, m in enumerate(self.months):
            put(ws, FF_CLASS_ROW + 1 + i, FF_MONTH_COL, month_text(m))
        if self.months:
            grid(ws, FF_CLASS_ROW + 1, FF_CLASS_ROW + len(self.months), 1, col=FF_MONTH_COL)
            band(ws, FF_CLASS_ROW + 1, FF_CLASS_ROW + len(self.months), FF_MONTH_COL, 1)
            align(ws, FF_CLASS_ROW + 1, FF_CLASS_ROW + len(self.months), 'L', FF_MONTH_COL)
        widths(ws, FF_WIDTHS)
        finish_fonts(ws)

    def order_intake(self, ws):
        o, rows = self.opts, self.data.oi
        n_files = sum(f.usable for f in self.kind_files('OI'))
        if not o.show_cancelled and not o.show_ssl:
            hidden = 'Filtered out: cancelled orders, skid-steer loaders.'
        elif not o.show_cancelled:
            hidden = 'Filtered out: cancelled orders.'
        elif not o.show_ssl:
            hidden = 'Filtered out: skid-steer loaders.'
        else:
            hidden = 'No rows filtered out.'
        title(ws, 'Order Intake', f'Source: Snowflake Order Intake export, {plural(n_files, "file")}, '
                                  f'built {self.built:%Y-%m-%d}. {hidden}', 11, OI_NOTE_COL)
        c = put(ws, 1, 6, OI_TOTAL_LABEL)                # the visible total, in the frozen title row
        c.font, c.alignment = BOLD_FONT, Alignment(horizontal='right')
        c = ws.cell(row=1, column=7, value=f'=SUBTOTAL(109,$G$3:$G${max(2 + len(rows), 3)})')   # skips hidden rows
        c.font, c.number_format, c.alignment = OI_TOTAL_FONT, UNITS_FMT, Alignment(horizontal='left', indent=1)
        header_row(ws, 2, 1, OI_HEADERS, OI_ALIGN)
        groups = set()
        for i, r in enumerate(rows):
            row = 3 + i
            if i == 0 or r['cls'] != rows[i - 1]['cls']:
                groups.add(row)
            for col, v in enumerate([r['cls'], self.data.registry[r['city']]] + r['vals'], start=1):
                c = put(ws, row, col, v)
                if col in (3, 4):
                    c.number_format = DATE_FMT
                elif col == 7:
                    c.number_format = QTY_FMT
            if r['cls'] == UNMAPPED:
                ws.cell(row=row, column=1).font = BOLD_FONT
            if r['hidden']:
                ws.row_dimensions[row].hidden = True
        if rows:                                       # UNMAPPED peach, first rule so it wins over the band
            ws.conditional_formatting.add(f'A3:A{2 + len(rows)}', FormulaRule(
                formula=['$A3="UNMAPPED"'], fill=FLAG_FILL, stopIfTrue=True))
        ruled_table(ws, 2, 2 + len(rows), OI_ALIGN, groups)
        ws.freeze_panes = 'A3'
        ws.auto_filter.ref = f'A2:K{2 + len(rows)}'
        if rows:
            if not o.show_cancelled:
                ws.auto_filter.filterColumn.append(FilterColumn(colId=4, customFilters=CustomFilters(
                    customFilter=[CustomFilter(operator='notEqual', val='CANCELLED')])))
            if not o.show_ssl:
                ws.auto_filter.add_filter_column(7, ['CTL', 'MX'])
        widths(ws, OI_WIDTHS)
        finish_fonts(ws)

    def inventory(self, ws):
        o, rows, n = self.opts, self.data.inv, len(self.cities)
        mtext = month_text(self.month)
        n_files = sum(f.usable for f in self.kind_files('INV'))
        detail = (f' Detail filtered to {mtext}{"" if o.show_ssl else ", CTL and MX"}.' if mtext else ' No inventory rows.')
        title(ws, 'Inventory by Month',
              f'Source: Snowflake Inventory by Month export, {plural(n_files, "file")}, built {self.built:%Y-%m-%d}. '
              f'Units: sum of TOTAL for the month in B3.{detail}', len(INV_HEADERS))
        ws['A3'] = 'Units in'
        ws['A3'].font = BOLD_FONT
        c = put(ws, 3, 2, mtext or None)
        c.number_format = '@'                          # Text: a month picked in the drop-down stays text, not a date
        c.border = Border(left=SIDES['thin'], right=SIDES['thin'], top=SIDES['thin'], bottom=SIDES['thin'])
        c.fill = INPUT_FILL
        align(ws, 3, 3, 'L', 2)
        if self.months:
            dv = DataValidation(type='list', formula1=f"'File Format'!$F${FF_CLASS_ROW + 1}:$F${FF_CLASS_ROW + len(self.months)}",
                                allow_blank=True)
            ws.add_data_validation(dv)
            dv.add('B3')
        header_row(ws, 4, 1, SUM_HEADERS, SUM_ALIGN)
        h = self.inv_head = 7 + n
        first, last = h + 1, h + max(len(rows), 1)
        rng = lambda col: f'${col}${first}:${col}${last}'
        for i, city in enumerate(self.cities):
            r = 5 + i
            put(ws, r, 1, city)
            for j, tname in enumerate(TYPE_NAMES):
                ws.cell(row=r, column=2 + j, value=f'=SUMIFS({rng("G")},{rng("A")},$A{r},{rng("B")},"{tname}",{rng("F")},$B$3)')
            ws.cell(row=r, column=5, value=f'=SUM(B{r}:D{r})')
            for col in range(2, 6):
                ws.cell(row=r, column=col).number_format = UNITS_FMT
        tr = 5 + n
        ws.cell(row=tr, column=1, value='Total').font = BOLD_FONT
        for col in range(2, 6):
            letter = openpyxl.utils.get_column_letter(col)
            c = ws.cell(row=tr, column=col, value=f'=SUM({letter}5:{letter}{4 + n})' if n else 0)
            c.font, c.number_format = BOLD_FONT, UNITS_FMT
        ruled_table(ws, 4, tr, SUM_ALIGN, unbanded=1)
        grid(ws, tr, tr, 5, {tr: 'double', tr + 1: 'thin'})         # the Total row: a double rule above, not banded
        header_row(ws, h, 1, INV_HEADERS, INV_ALIGN)
        groups = set()
        for i, r in enumerate(rows):
            row = first + i
            if i == 0 or r['city'] != rows[i - 1]['city']:
                groups.add(row)
            v = r['vals']
            for col, value in enumerate([self.data.registry[r['city']], v[0], v[1], v[2], v[3], month_text(r['month']) or None,
                                         v[4], v[5], v[6]], start=1):
                c = put(ws, row, col, value)
                if col == 5:
                    c.number_format = DATE_FMT
                elif col == 7:
                    c.number_format = UNITS_FMT
            if r['hidden']:
                ws.row_dimensions[row].hidden = True
        ruled_table(ws, h, h + len(rows), INV_ALIGN, groups)
        ws.freeze_panes = f'A{first}'
        ws.auto_filter.ref = f'A{h}:I{h + len(rows)}'
        if rows:
            if self.month is not None:
                ws.auto_filter.add_filter_column(5, [mtext])
            if not o.show_ssl:
                ws.auto_filter.add_filter_column(1, TYPE_NAMES[:2])
        widths(ws, INV_WIDTHS)
        finish_fonts(ws)

    def open_orders(self, ws):
        """Open Orders by MSA: the ordered quantity of the open Order Intake rows by MSA (alphabetical, every MSA of
        the report) and product type, as SUMIFS over the Order Intake tab, a Total column and row, then the chart."""
        o, n = self.opts, len(self.cities)
        nt = 3 if o.show_ssl else 2
        title(ws, OO_SHEET, f'Source: the Order Intake tab, built {self.built:%Y-%m-%d}. Ordered quantity of orders with '
                            f'ORDER_STATUS ORDERED or Not yet processed, by MSA and product type.'
                            + (' Chart: CTL, MX and SSL.' if o.show_ssl else ' Chart: CTL and MX.'), 5)
        caption(ws, 3, 1, 'Table 1.', 'Open orders by MSA: sum of SUM(F.ORDERED_QUANTITY) over the Order Intake rows whose '
                                      'ORDER_STATUS is ORDERED or Not yet processed.')
        header_row(ws, 4, 1, SUM_HEADERS, SUM_ALIGN)
        last = 2 + max(len(self.data.oi), 1)
        rng = lambda col: f"'Order Intake'!${col}$3:${col}${last}"
        for i, city in enumerate(self.cities):
            r = 5 + i
            put(ws, r, 1, city)
            for j in range(3):
                L = openpyxl.utils.get_column_letter(2 + j)
                ws.cell(row=r, column=2 + j, value='=' + '+'.join(
                    f'SUMIFS({rng("G")},{rng("B")},$A{r},{rng("H")},{L}$4,{rng("E")},"{st}")' for st in OPEN_STATUSES))
            ws.cell(row=r, column=5, value=f'=SUM(B{r}:D{r})')
            for col in range(2, 6):
                ws.cell(row=r, column=col).number_format = UNITS_FMT
        tr = 5 + n
        ws.cell(row=tr, column=1, value='Total').font = BOLD_FONT
        for col in range(2, 6):
            letter = openpyxl.utils.get_column_letter(col)
            c = ws.cell(row=tr, column=col, value=f'=SUM({letter}5:{letter}{4 + n})' if n else 0)
            c.font, c.number_format = BOLD_FONT, UNITS_FMT
        ruled_table(ws, 4, tr, SUM_ALIGN, unbanded=1)
        grid(ws, tr, tr, 5, {tr: 'double', tr + 1: 'thin'})         # the Total row: a double rule above, not banded
        widths(ws, OO_WIDTHS)
        if n:
            add_open_orders_chart(ws, n, nt, 7 + n)
        else:
            note(ws, 7 + n, 1, 'No MSAs: no chart.')
        finish_fonts(ws)

    def run_log(self, ws):
        secs = time.time() - self.started
        put(ws, 1, 1, f'Built {self.built:%Y-%m-%d %H:%M} from {os.path.abspath(self.folder)} in {secs:.1f} s')
        put(ws, 2, 1, self.options_line())
        for i, line in enumerate(self.log, start=4):
            put(ws, i, 1, line)
        ws.column_dimensions['A'].width = LOG_WIDTH
        finish_fonts(ws, LOG_FONT)

def parse_args(argv=None):
    ap = argparse.ArgumentParser(description='Build the Open Orders report (v2) without the macro.')
    ap.add_argument('input_folder')
    ap.add_argument('output_xlsx')
    ap.add_argument('--size-class-file', required=True, help='the size class workbook (master list)')
    ap.add_argument('--msa-column', dest='city_column', default='',
                    help='header of the MSA column (default: found automatically, an MSA header first)')
    ap.add_argument('--inventory-month', default='', help='any day in the inventory month to show, YYYY-MM-DD (default: latest)')
    ap.add_argument('--show-cancelled', action='store_true', help='do not hide CANCELLED orders behind the filter')
    ap.add_argument('--show-ssl', action='store_true', help='do not hide skid-steer loaders behind the filters')
    opts = ap.parse_args(argv)
    opts.inventory_month = dt.datetime.strptime(opts.inventory_month, '%Y-%m-%d').date() if opts.inventory_month else None
    return opts

def main(argv=None):
    opts = parse_args(argv)
    b = Build(opts.input_folder, opts)
    if b.problems:
        print('The report was not built. Fix this:', file=sys.stderr)
        for p in b.problems:
            print('PROBLEM: ' + p, file=sys.stderr)
        return 2
    b.write(opts.output_xlsx)
    print('\n'.join(b.log))
    print('Saved', opts.output_xlsx)
    return 0

if __name__ == '__main__':
    sys.exit(main())
