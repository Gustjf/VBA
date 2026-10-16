#!/usr/bin/env python3
"""
Write the screenshots of the size-class master workbook used by size-class-file-guide.tex:
source/open-orders-intake-inventory/size-class-guide/*.png

Each picture is the real workbook (output/open-orders-intake-inventory/size-class-master.xlsx) printed by LibreOffice
with its row and column headings, one range per picture, cropped to the content. The "after" pictures add the guide's
worked example (CTL 4000&Over, SSL XL Frame, MX >10-15, model E145) in a copy; the cells typed are shaded peach and
bold for the guide only. The workbook itself is never changed.
Run it again, then ./build.sh the .tex, whenever the workbook's look or its classes change.

Needs: soffice (LibreOffice), pdftoppm (poppler-utils), openpyxl, Pillow.
Usage: python3 make_size_class_guide_shots.py
"""
import copy, os, subprocess, tempfile
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties
from PIL import Image, ImageChops

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', '..', 'output', 'open-orders-intake-inventory', 'size-class-master.xlsx')
OUT = os.path.join(HERE, 'size-class-guide')
PEACH = PatternFill('solid', fgColor='FCE4D6')
# the guide's worked example: (sheet, cell, value, neighbour whose border and font it takes)
NEW_CLASSES = [('Set-up', 'A11', '4000&Over', 'A10'), ('Set-up', 'C9', 'XL Frame', 'C8'), ('Set-up', 'E12', '>10-15', 'E11')]
NEW_MODEL = [('Size Class', 'C31', 'E145', 'C30'), ('Size Class', 'D31', '>10-15', 'D30'),
             ('Size Class', 'E31', 'new model, new class', 'E30')]

def shot(tmp, name, sheet, area, edits=(), widths=None, hide=()):
    wb = openpyxl.load_workbook(SRC)
    for sh, cell, value, like in edits:
        c, n = wb[sh][cell], wb[sh][like]
        c.value = value
        c.border, c.alignment = copy.copy(n.border), copy.copy(n.alignment)
        c.fill, c.font = PEACH, Font(name=n.font.name, size=n.font.sz, bold=True)
    ws = wb[sheet]
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = w
    for r in hide:
        ws.row_dimensions[r].hidden = True
    for other in wb:
        if other.title != sheet:
            other.sheet_state = 'hidden'
    wb.active = wb.sheetnames.index(sheet)
    ws.print_area = area
    ws.print_options.headings = True
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    xlsx = os.path.join(tmp, name + '.xlsx')
    wb.save(xlsx)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, xlsx], check=True, capture_output=True)
    subprocess.run(['pdftoppm', '-png', '-r', '220', '-singlefile', xlsx[:-5] + '.pdf', os.path.join(tmp, name)], check=True)
    im = Image.open(os.path.join(tmp, name + '.png')).convert('RGB')
    box = ImageChops.difference(im, Image.new('RGB', im.size, 'white')).getbbox()
    im = im.crop((box[0] - 4, box[1] - 4, box[2] + 4, box[3] + 4))
    im.save(os.path.join(OUT, name + '.png'))
    print(f'{name}.png {im.size[0]} x {im.size[1]}')

def main():
    os.makedirs(OUT, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        shot(tmp, 'setup-tab', 'Set-up', 'A1:I22')
        shot(tmp, 'size-class-tab', 'Size Class', 'A1:E30', widths={'E': 80})
        shot(tmp, 'setup-after', 'Set-up', 'A3:G24', edits=NEW_CLASSES)
        shot(tmp, 'size-class-new', 'Size Class', 'A1:E31', widths={'E': 40}, hide=range(2, 26),
             edits=NEW_CLASSES + NEW_MODEL)

if __name__ == '__main__':
    main()
