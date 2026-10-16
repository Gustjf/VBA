#!/usr/bin/env python3
"""
Compare a v2 report built from sample-data (macro or build_report_reference.py, InventoryMonth 2026-11-08) with
the hand-built Austin tab in sample-data.

Usage:
    python3 compare_with_handbuilt.py <report.xlsx> [--city Austin] [--handbuilt <hand-built.xlsx>]

Expected result (plan Section 7.3):
    Order Intake  the visible rows (not hidden) of the MSA: 170 of 170 rows equal to the hand-built left block as
                  a set (columns ORDER_DATE..ORGANIZATION_NAME), the Size Class column identical in order. Rows may
                  now sit elsewhere inside a class because of the order status key: the count in the same position
                  is reported, not failed.
    Inventory     the MSA's rows of the v2 detail block for the hand-built month (November 2026), CTL and MX:
                  29 of 29 equal in order (PRODUCT_TYPE, MODEL_GROUP, MODEL_NUMBER, INVENTORY_DATE), at most 3
                  TOTAL values different (anonymisation noise).
    Summary       the CTL / MX cells are not compared with the hand-built 21 / 2 (anonymised sample).
Exit code 0 when all of that holds, 1 otherwise.
"""
import argparse, collections, os, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
HANDBUILT = os.path.join(HERE, 'sample-data', '001. Open Orders - Order Intake - Inventory Analysis By MSA_ANONYMIZED_TRIMMED.xlsx')
MONTH = 'November 2026'


def block(ws, c0, c1, r0):
    out, r = [], r0
    while ws.cell(r, c0).value not in (None, ''):
        out.append(tuple(ws.cell(r, c).value for c in range(c0, c1 + 1)))
        r += 1
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('report')
    ap.add_argument('--city', default='Austin')
    ap.add_argument('--handbuilt', default=HANDBUILT)
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.report)
    oi, inv = wb['Order Intake'], wb['Inventory by Month']
    hand = openpyxl.load_workbook(a.handbuilt)['Austin']
    results = []

    def check(name, ok, detail):
        results.append(ok)
        print(f"{'PASS' if ok else 'FAIL'} {name}: {detail}")

    built = [tuple(oi.cell(r, c).value for c in [1] + list(range(3, 12))) for r in range(3, oi.max_row + 1)
             if oi.cell(r, 2).value == a.city and not oi.row_dimensions[r].hidden]
    hl = block(hand, 1, 10, 3)
    same_pos = sum(x == y for x, y in zip(built, hl))
    check('visible rows as a set', collections.Counter(r[1:] for r in built) == collections.Counter(r[1:] for r in hl),
          f'{len(built)} visible {a.city} rows built, {len(hl)} hand-built')
    check('size-class order', [r[0] for r in built] == [r[0] for r in hl], 'Size Class column in the same order')
    print(f'INFO rows in the same position: {same_pos} of {len(hl)} (not a check: the status key moves rows inside a class)')

    h = next(r for r in range(5, inv.max_row + 1) if inv.cell(r, 1).value == 'MSA' and inv.cell(r, 2).value == 'PRODUCT_TYPE')
    bi = [tuple(inv.cell(r, c).value for c in (2, 3, 4, 5, 7)) for r in range(h + 1, inv.max_row + 1)
          if inv.cell(r, 1).value == a.city and inv.cell(r, 6).value == MONTH
          and inv.cell(r, 2).value in ('Compact Track Loaders', 'Mini Excavators')]
    hi = block(hand, 17, 21, 8)
    check('inventory rows', len(bi) == len(hi) and [r[:4] for r in bi] == [r[:4] for r in hi],
          f'{len(bi)} built, {len(hi)} hand-built, PRODUCT_TYPE..INVENTORY_DATE equal in order')
    diffs = [(i + 1, x[4], y[4]) for i, (x, y) in enumerate(zip(bi, hi)) if x[4] != y[4]]
    check('inventory TOTAL', len(diffs) <= 3, f'{len(diffs)} differ (row, built, hand-built): {diffs}')
    print(f"RESULT: {'PASS' if all(results) else 'FAIL'} ({sum(results)} of {len(results)} checks)")
    sys.exit(0 if all(results) else 1)


if __name__ == '__main__':
    main()
