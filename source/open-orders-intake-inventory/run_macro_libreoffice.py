#!/usr/bin/env python3
"""
Run the VBA macro (build-open-orders-report.bas, version 2.8) in headless LibreOffice Calc, so its logic can be
executed on a machine without Excel (for example a cloud session), then check the result with check_report.py
like any other report.

This is extra evidence, not the Excel acceptance test: LibreOffice runs VBA in a compatibility mode that lacks
a few Excel features, so the module is loaded with the test-only substitutions below and the result is
normalised afterwards. Everything else (discovery, reading, the MSA column, the size class file, sorting,
text dates, the five tabs, formulas, filters, Run Log, Config, the setup check, saving) runs as written; of the
Open Orders by MSA tab only the chart is substituted (its title, table, formulas, rules and page setup run as written);
the Order Intake visible total (v2.6, F1:G1) runs as written.

    Substitution (test copy only)                     Why
    MsgBox -> lines in <output>.msg.log               headless LibreOffice cannot show dialogs
    EnsureButtons call skipped                        no Shapes.AddFormControl in LibreOffice
    Date and hint-only Config checks skipped          LibreOffice's Validation.Add has only list and whole number
    Protect without the Allow... arguments            LibreOffice's Worksheet.Protect does not take them
    Application.ScreenUpdating = False skipped        with it, ws.Cells fails after Workbooks.Open/Close
    FileFormat:=xlOpenXMLWorkbook -> 51               LibreOffice does not define the constant
    ActiveWindow.DisplayGridlines = False             LibreOffice keeps no per-sheet gridline setting from VBA; the
      -> TestGridlines ws (records the sheet)         normaliser turns gridlines off on the recorded sheets
    ActiveWindow.SplitRow / FreezePanes               LibreOffice froze Order Intake at A6 and Inventory not at all;
      -> TestFreeze ws, firstRow (records the row)    the normaliser freezes the recorded sheet above the row
    rng.AutoFilter Field:=n, Criteria1:=...           LibreOffice's VBA AutoFilter keeps one criterion and rejects
      -> TestAutoFilter rng, n, op, values            xlFilterValues; the harness hides exactly the rows the
                                                      criterion excludes (Excel's rule: ignore case, compare the
                                                      whole text) and records it; the normaliser writes the filter
                                                      column (a value list, or "not equal" for <>CANCELLED)
    Validation Formula1 "='File Format'!$F$4:..."     LibreOffice misreads the Excel-style sheet reference here;
      -> "$'File Format'.$F$4:..."                    its own syntax gives the same range in the saved file
    SetRule (the medium and double rules)             LibreOffice's VBA ignores xlDouble and saves xlMedium as
      -> TestRule rng, edge, style (records it)       thick; the normaliser draws the recorded rules
    rng.FormatConditions.Add in AddFillRule           LibreOffice's VBA cannot give a conditional format a fill;
      (row banding, UNMAPPED peach)                   the normaliser writes the recorded formula and fill, the
      -> TestCF rng, f, fill, first (records it)      rules marked first ahead of the others
    ws.PageSetup properties in SetUpPage              LibreOffice does not read Excel's header and footer codes
      -> TestPage ws, name, value (records each)      (&A, &P, the font); the normaliser writes the recorded page
                                                      setup (orientation, fit, margins, header, footer, titles)
    DrawOpenOrdersChart ws, nc, nt, anchorRow         LibreOffice's VBA chart support is partial (and its own chart
      -> TestChart ws, nc, nt, anchorRow              would not be Excel's), so the call is recorded and the
      (records it)                                    normaliser draws the chart, once, after loading the file, with
                                                      the reference's own add_open_orders_chart
                                                      (build_report_reference.py)
    Output normalised: booleans written as =TRUE()/=FALSE() by LibreOffice's xlsx export become TRUE/FALSE
    again; the recorded gridlines, freeze panes, filter columns, rules, conditional fills and page setup are written (rows the
    harness hid stay hidden); LibreOffice saves .xls content under the .xlsx name, so it is converted.

Usage (same options as build_report_reference.py, typed into Config the way a person would):
    python3 run_macro_libreoffice.py <input_folder> <output.xlsx> --size-class-file PATH [--msa-column NAME]
                                     [--inventory-month YYYY-MM-DD] [--show-cancelled] [--show-ssl]
                                     [--month-as-text] [--check-only] [--set KEY=VALUE ...] [--v1-config | --v20-config]
                                     [--bas PATH]
    InventoryMonth goes into Config as a real date (as when a person types 8/15/2026), or as text with
    --month-as-text. --set types a value into any Config setting afterwards (for example a bad one, to see the
    setup check stop the build); --check-only clicks Check setup instead of Build report. --v1-config first
    writes a version 1.3 Config sheet (and a SizeClassMap sheet), --v20-config a version 2.0 one (CityColumn),
    then runs ShowOpenOrdersReportConfig (the upgrade) and Check setup, and prints the upgraded settings. The messages and the Config "Last run" box are
    printed after every run.
Then, for example:
    python3 check_report.py <output.xlsx> --expected test-fixtures/expected.json
    python3 check_report.py <output.xlsx> --compare <same build from build_report_reference.py>
Needs LibreOffice (soffice) and its Python UNO bindings (python3 -c "import uno").
"""
import argparse, os, re, sys, time, shutil, socket, tempfile, subprocess

HERE = os.path.dirname(os.path.abspath(__file__))
# (old text, new text, count, why) for the substitutions after the first six in the docstring table; each must
# match exactly count times, so a change to those lines in the module stops the test instead of passing silently.
EXTRA = [
    ("""    ws.Activate
    ActiveWindow.DisplayGridlines = False
""", """    TestGridlines ws
""", 1, 'gridlines off'),
    ("""    ws.Activate
    With ActiveWindow
        .FreezePanes = False
        .ScrollRow = 1
        .ScrollColumn = 1
        .SplitColumn = 0
        .SplitRow = firstRow - 1
        .FreezePanes = True
    End With
    ws.Range("A1").Select
""", """    TestFreeze ws, firstRow
""", 1, 'freeze panes'),
    ('Formula1:="=\'File Format\'!$F$"', 'Formula1:="$\'File Format\'.$F$"', 1, 'validation list reference'),
    ("""Private Sub SetRule(rng As Range, ByVal edge As Long, ByVal style As String)
    With rng.Borders(edge)
        If style = "double" Then
            .LineStyle = xlDouble
        Else
            .LineStyle = xlContinuous
            .Weight = IIf(style = "medium", xlMedium, xlThin)
        End If
        .Color = 0
    End With
End Sub
""", """Private Sub SetRule(rng As Range, ByVal edge As Long, ByVal style As String)
    TestRule rng, edge, style
End Sub
""", 1, 'medium and double rules'),
    ("""    rng.Worksheet.Activate                               ' Excel reads the formula relative to the active cell
    rng.Cells(1, 1).Select
    With rng.FormatConditions.Add(Type:=xlExpression, Formula1:=f)
        .Interior.Color = fillColor
        .StopIfTrue = first
        If first Then .SetFirstPriority
    End With
""", """    TestCF rng, f, fillColor, first
""", 1, 'conditional fills'),
    ('        DrawOpenOrdersChart ws, nc, nt, anchorRow\n', '        TestChart ws, nc, nt, anchorRow\n', 1, 'open orders chart'),
]
# AutoFilter criteria calls "rng.AutoFilter Field:=n, Criteria1:=..." -> "TestAutoFilter rng, n, op, Array(...)"
AUTOFILTER_CALLS = 4

def free_port():
    s = socket.socket(); s.bind(('127.0.0.1', 0)); port = s.getsockname()[1]; s.close()
    return port

def test_copy(bas_path, msg_log):
    """The module text with the LibreOffice-only substitutions, and the harness module."""
    src = open(bas_path, 'rb').read().decode('ascii').replace('\r\n', '\n').split('\n')
    if not src[0].startswith('Attribute VB_Name'):
        sys.exit(f'{bas_path}: first line is not "Attribute VB_Name"')
    code = 'Option VBASupport 1\n' + '\n'.join(src[1:])
    code = re.sub(r'\bMsgBox\(', 'TestMsgFn(', code)
    code = re.sub(r'\bMsgBox ', 'TestMsg ', code)
    subs = [('\n    EnsureButtons ws\n', "\n    ' EnsureButtons skipped in LibreOffice\n"),
            ('    Application.ScreenUpdating = False\n', "    ' Application.ScreenUpdating = False skipped in LibreOffice\n"),
            ('FileFormat:=xlOpenXMLWorkbook', 'FileFormat:=51'),
            ('    If c Is Nothing Then Exit Sub\n    With c.Validation\n',
             '    If c Is Nothing Then Exit Sub\n    If checkType <> xlValidateList And checkType <> xlValidateWholeNumber Then Exit Sub\n'
             '    With c.Validation\n')]
    for old, new, count in [(o, n, 1) for o, n in subs] + [(o, n, c) for o, n, c, _ in EXTRA]:
        if code.count(old) != count:
            sys.exit(f'expected {count} x {old.strip()[:70]!r} in the module, found {code.count(old)}; '
                     f'update run_macro_libreoffice.py')
        code = code.replace(old, new)
    a = code.index('Private Sub SetUpPage(')
    b = code.index('End Sub', a)
    body, k = re.subn(r'\.(\w+) = ', r'TestPage ws, "\1", ', code[a:b])
    body = body.replace('    With ws.PageSetup\n', '').replace('    End With\n', '')
    if k != 16 or 'PageSetup' in body:
        sys.exit(f'expected 16 page setup property assignments in SetUpPage, found {k}; update run_macro_libreoffice.py')
    code = code[:a] + body + code[b:]
    def autofilter(m):
        field, crit = m.group(1), m.group(2)
        if crit.startswith('"<>'):
            return f'TestAutoFilter rng, {field}, "<>", Array("{crit[3:-1]}")'
        arr = re.fullmatch(r'(Array\(.*\)), Operator:=xlFilterValues', crit)
        if not arr:
            sys.exit(f'AutoFilter criteria the harness does not know: {crit}; update run_macro_libreoffice.py')
        return f'TestAutoFilter rng, {field}, "=", {arr.group(1)}'
    code, n = re.subn(r'rng\.AutoFilter Field:=(\d+), Criteria1:=([^\n]+)', autofilter, code)
    if n != AUTOFILTER_CALLS:
        sys.exit(f'expected {AUTOFILTER_CALLS} AutoFilter criteria calls, found {n}; update run_macro_libreoffice.py')
    code, n = re.subn(r'(ws\.Protect DrawingObjects:=True, Contents:=True, Scenarios:=True)(?:[^\n]*_\n)*[^\n]*', r'\1', code)
    if n != 1:
        sys.exit(f'expected one ws.Protect call in the module, found {n}; update run_macro_libreoffice.py')
    if re.search(r'\.ThemeColor\s*=', code):
        sys.exit('the module sets a ThemeColor; the report uses no theme colours')
    q = lambda s: s.replace('"', '""')
    status_row = int(re.search(r'Private Const STATUS_ROW As Long = (\d+)', code).group(1))
    harness = f'''Option VBASupport 1
Public Sub TestMsg(ByVal prompt As String, Optional ByVal buttons As Long = 0, Optional ByVal title As String = "")
    Dim f As Integer
    f = FreeFile
    Open "{q(msg_log)}" For Append As #f
    Print #f, "MSGBOX: " & prompt
    Close #f
End Sub
Public Function TestMsgFn(ByVal prompt As String, Optional ByVal buttons As Long = 0, Optional ByVal title As String = "") As Long
    TestMsg prompt, buttons, title
    TestMsgFn = 6
End Function
Private Sub Record(ByVal s As String)
    Dim f As Integer
    f = FreeFile
    Open "{q(msg_log)}" For Append As #f
    Print #f, "RECORD|" & s
    Close #f
End Sub
Public Sub TestGridlines(ws As Object)
    Record "GRIDLINES|" & ws.Name
End Sub
Public Sub TestFreeze(ws As Object, ByVal firstRow As Long)
    Record "FREEZE|" & ws.Name & "|" & firstRow
End Sub
Public Sub TestRule(rng As Object, ByVal edge As Long, ByVal style As String)
    Record "RULE|" & rng.Worksheet.Name & "|" & rng.Row & "|" & rng.Column & "|" & (rng.Row + rng.Rows.Count - 1) & "|" & _
           (rng.Column + rng.Columns.Count - 1) & "|" & edge & "|" & style
End Sub
Public Sub TestCF(rng As Object, ByVal f As String, ByVal fillColor As Long, ByVal first As Boolean)
    Record "CF|" & rng.Worksheet.Name & "|" & rng.Row & "|" & rng.Column & "|" & (rng.Row + rng.Rows.Count - 1) & "|" & _
           (rng.Column + rng.Columns.Count - 1) & "|" & fillColor & "|" & IIf(first, "1", "0") & "|" & f
End Sub
Public Sub TestPage(ws As Object, ByVal prop As String, ByVal v As Variant)
    Record "PAGE|" & ws.Name & "|" & prop & "|" & CStr(v)
End Sub
Public Sub TestChart(ws As Object, ByVal nc As Long, ByVal nt As Long, ByVal anchorRow As Long)
    Record "CHART|" & ws.Name & "|" & nc & "|" & nt & "|" & anchorRow
End Sub
' Excel's AutoFilter rules: "<>v" hides the rows whose text is v (ignoring case); a value list keeps only the rows
' whose text is one of the values (ignoring case). Rows already hidden by another column stay hidden.
Public Sub TestAutoFilter(rng As Object, ByVal field As Long, ByVal op As String, ByVal vals As Variant)
    Dim r As Long, t As String, i As Long, inList As Boolean, s As String
    For r = 2 To rng.Rows.Count
        t = UCase$(CStr(rng.Cells(r, field).Value))
        inList = False
        For i = LBound(vals) To UBound(vals)
            If t = UCase$(CStr(vals(i))) Then inList = True
        Next i
        If (op = "<>" And inList) Or (op = "=" And Not inList) Then rng.Rows(r).EntireRow.Hidden = True
    Next r
    For i = LBound(vals) To UBound(vals)
        s = s & Chr(31) & vals(i)
    Next i
    Record "AUTOFILTER|" & rng.Worksheet.Name & "|" & field & "|" & op & "|" & Mid$(s, 2)
End Sub
Private Sub SetByKey(ByVal k As String, ByVal v As Variant)
    Dim r As Long
    With ThisWorkbook.Worksheets("Config")
        For r = 2 To 10
            If .Cells(r, 1).Value = k Then .Cells(r, 2).Value = v
        Next r
    End With
End Sub
Private Sub PrintLastRun()
    Dim r As Long
    With ThisWorkbook.Worksheets("Config")
        For r = {status_row} To {status_row} + 5
            TestMsg "LAST RUN " & .Cells(r, 1).Value & ": " & .Cells(r, 2).Text
        Next r
    End With
End Sub
Private Sub WriteOldConfig(ByVal version As String, ByVal inFolder As String, ByVal outFolder As String, ByVal outName As String, _
                           ByVal sizeFile As String)
    Dim ws As Object, keys As Variant, vals As Variant, i As Long
    Set ws = ThisWorkbook.Worksheets.Add
    ws.Name = "Config"
    ws.Range("A1").Value = "Setting": ws.Range("B1").Value = "Value": ws.Range("C1").Value = "What it does"
    If version = "1.3" Then
        keys = Array("InputFolder", "OutputFolder", "OutputFileName", "InventorySnapshotDate", "DropCancelled", "IncludeSSL", _
                     "HeaderCountMode", "InventoryHeaderRow", "OpenAfterBuild")
        vals = Array(inFolder, outFolder, outName, DateSerial(2026, 11, 8), "TRUE", "FALSE", "InventoryUnits", "7", "FALSE")
    Else
        keys = Array("InputFolder", "OutputFolder", "OutputFileName", "SizeClassFile", "CityColumn", "InventoryMonth", _
                     "ShowCancelled", "ShowSSL", "OpenAfterBuild")
        vals = Array(inFolder, outFolder, outName, sizeFile, "MSA", DateSerial(2026, 8, 15), "TRUE", "FALSE", "FALSE")
    End If
    For i = 0 To 8
        ws.Cells(2 + i, 1).Value = keys(i)
        ws.Cells(2 + i, 2).Value = vals(i)
        ws.Cells(2 + i, 3).Value = "(version " & version & " help text)"
    Next i
    If version = "1.3" Then
        Set ws = ThisWorkbook.Worksheets.Add
        ws.Name = "SizeClassMap"
        ws.Range("A1").Value = "MODEL_GROUP": ws.Range("C1").Value = "SIZE_CLASS"
    End If
End Sub
Public Sub RunTest(ByVal inFolder As String, ByVal outFolder As String, ByVal outName As String, ByVal sizeFile As String, _
                   ByVal cityCol As String, ByVal invMonth As String, ByVal showC As String, ByVal showS As String, _
                   ByVal monthAsText As Boolean, ByVal checkOnly As Boolean, ByVal overrides As String, ByVal oldConfig As String)
    Dim kv As Variant, i As Long, r As Long
    On Error GoTo E
    If Len(oldConfig) > 0 Then
        WriteOldConfig oldConfig, inFolder, outFolder, outName, sizeFile
        ShowOpenOrdersReportConfig
        With ThisWorkbook.Worksheets("Config")
            For r = 1 To 10
                TestMsg "CONFIG ROW " & r & ": " & .Cells(r, 1).Value & " = " & .Cells(r, 2).Text
            Next r
        End With
        PrintLastRun
        CheckOpenOrdersReportSetup
        PrintLastRun
        Exit Sub
    End If
    ShowOpenOrdersReportConfig
    SetByKey "InputFolder", inFolder
    SetByKey "OutputFolder", outFolder
    SetByKey "OutputFileName", outName
    SetByKey "SizeClassFile", sizeFile
    SetByKey "MSAColumn", cityCol
    If Len(invMonth) = 10 And Not monthAsText Then          ' a typed date, as Excel stores it
        SetByKey "InventoryMonth", DateSerial(CLng(Left$(invMonth, 4)), CLng(Mid$(invMonth, 6, 2)), CLng(Right$(invMonth, 2)))
    Else
        SetByKey "InventoryMonth", invMonth
    End If
    SetByKey "ShowCancelled", showC
    SetByKey "ShowSSL", showS
    SetByKey "OpenAfterBuild", "FALSE"
    If Len(overrides) > 0 Then                         ' --set KEY=VALUE, typed into that setting's Value cell
        kv = Split(overrides, Chr(30))
        For i = LBound(kv) To UBound(kv)
            SetByKey Left$(kv(i), InStr(kv(i), "=") - 1), Mid$(kv(i), InStr(kv(i), "=") + 1)
        Next i
    End If
    If checkOnly Then CheckOpenOrdersReportSetup Else BuildOpenOrdersReport
    PrintLastRun
    Exit Sub
E:
    TestMsg "HARNESS ERROR " & Err.Number & ": " & Err.Description
End Sub
'''
    return code, harness

XL_EDGE_TOP, XL_EDGE_BOTTOM, XL_LANDSCAPE = 8, 9, 2

def draw_rule(ws, r1, c1, r2, c2, edge, style):
    """The rule Excel draws for rng.Borders(edge): on the edge cells and on their neighbours across the edge."""
    from openpyxl.styles import Border, Side
    side = Side(style=style, color='FF000000')
    def put(r, c, which):
        if r < 1:
            return
        b = ws.cell(row=r, column=c).border
        parts = {k: getattr(b, k) for k in ('left', 'right', 'top', 'bottom')}
        parts[which] = side
        ws.cell(row=r, column=c).border = Border(**parts)
    for c in range(c1, c2 + 1):
        if edge == XL_EDGE_TOP:
            put(r1, c, 'top'); put(r1 - 1, c, 'bottom')
        elif edge == XL_EDGE_BOTTOM:
            put(r2, c, 'bottom'); put(r2 + 1, c, 'top')
        else:
            sys.exit(f'TestRule edge {edge} the harness does not draw; update run_macro_libreoffice.py')

def apply_page(ws, props):
    """The page setup Excel would store for the recorded PageSetup properties."""
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.worksheet.properties import PageSetupProperties
    for part in (ws.oddHeader.left, ws.oddHeader.center, ws.oddHeader.right,
                 ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
        part.text = part.font = part.size = None         # LibreOffice's own (mangled) header and footer go
    ps = ws.page_setup
    ps.orientation = 'landscape' if props.get('Orientation') == str(XL_LANDSCAPE) else 'portrait'
    if 'FitToPagesWide' in props:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ps.fitToWidth = int(props['FitToPagesWide'])
        tall = props.get('FitToPagesTall', '1')
        ps.fitToHeight = 0 if tall.lower() == 'false' else int(tall)
    pt = lambda k: float(props[k].replace(',', '.')) / 72
    ws.page_margins = PageMargins(left=pt('LeftMargin'), right=pt('RightMargin'), top=pt('TopMargin'),
                                  bottom=pt('BottomMargin'), header=pt('HeaderMargin'), footer=pt('FooterMargin'))
    parts = {'LeftHeader': ws.oddHeader.left, 'CenterHeader': ws.oddHeader.center, 'RightHeader': ws.oddHeader.right,
             'LeftFooter': ws.oddFooter.left, 'CenterFooter': ws.oddFooter.center, 'RightFooter': ws.oddFooter.right}
    for k, part in parts.items():
        if k in props:
            m = re.fullmatch(r'&"([^"]+)"&(\d+)(.*)', props[k], re.S)
            part.font, part.size, part.text = (m.group(1), int(m.group(2)), m.group(3)) if m else (None, None, props[k])
    if props.get('PrintTitleRows'):
        ws.print_title_rows = props['PrintTitleRows'].replace('$', '')

def add_fill_rules(ws, rules):
    """The recorded conditional fills, the rules marked first (SetFirstPriority, StopIfTrue) ahead of the others."""
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter as L
    for rest in sorted(rules, key=lambda x: x[5] != '1'):
        r1, c1, r2, c2, colour = (int(x) for x in rest[:5])
        rgb = f'FF{colour & 255:02X}{(colour >> 8) & 255:02X}{(colour >> 16) & 255:02X}'
        ws.conditional_formatting.add(f'{L(c1)}{r1}:{L(c2)}{r2}', FormulaRule(
            formula=['|'.join(rest[6:]).lstrip('=')], fill=PatternFill('solid', start_color=rgb, end_color=rgb),
            stopIfTrue=rest[5] == '1'))

def normalise(raw_path, out_path, uno_call, records):
    """LibreOffice output (.xls content) -> .xlsx with booleans restored and the recorded gridlines, freeze panes,
    filter columns, rules, page setup and chart written."""
    import openpyxl
    from openpyxl.worksheet.filters import FilterColumn, CustomFilters, CustomFilter
    sys.path.insert(0, HERE)
    from build_report_reference import add_open_orders_chart
    tmp = out_path + '.lo-convert.xlsx'
    uno_call('convert', raw_path, tmp)
    wb = openpyxl.load_workbook(tmp)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value in ('=TRUE()', '=FALSE()'):
                    c.value = (c.value == '=TRUE()')
        ws.auto_filter.filterColumn = []                 # LibreOffice's leftovers; the recorded ones follow
    pages, cfs, charts = {}, {}, []
    for rec in records:
        kind, sheet, *rest = rec.split('|')
        ws = wb[sheet]
        if kind == 'CHART':                      # drawn last: the sheet is loaded, so the chart is not dropped
            charts.append((ws, *(int(x) for x in rest)))
        elif kind == 'PAGE':
            pages.setdefault(sheet, {})[rest[0]] = '|'.join(rest[1:])
        elif kind == 'CF':
            cfs.setdefault(sheet, []).append(rest)
        elif kind == 'RULE':
            r1, c1, r2, c2, edge = (int(x) for x in rest[:5])
            draw_rule(ws, r1, c1, r2, c2, edge, rest[5])
        elif kind == 'GRIDLINES':
            ws.sheet_view.showGridLines = False
        elif kind == 'FREEZE':
            ws.freeze_panes = f'A{rest[0]}'
        elif kind == 'AUTOFILTER':
            field, op, vals = int(rest[0]), rest[1], rest[2].split(chr(31))
            if op == '<>':
                ws.auto_filter.filterColumn.append(FilterColumn(colId=field - 1, customFilters=CustomFilters(
                    customFilter=[CustomFilter(operator='notEqual', val=v) for v in vals])))
            else:
                ws.auto_filter.add_filter_column(field - 1, vals)
    for sheet, props in pages.items():
        apply_page(wb[sheet], props)
    for sheet, rules in cfs.items():
        add_fill_rules(wb[sheet], rules)
    for ws, nc, nt, anchor_row in charts:
        add_open_orders_chart(ws, nc, nt, anchor_row)
    wb.save(out_path)
    os.remove(tmp)

def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('input_folder'); ap.add_argument('output_xlsx')
    ap.add_argument('--size-class-file', default='')
    ap.add_argument('--msa-column', dest='city_column', default='')
    ap.add_argument('--inventory-month', default='')
    ap.add_argument('--show-cancelled', action='store_true')
    ap.add_argument('--show-ssl', action='store_true')
    ap.add_argument('--month-as-text', action='store_true', help='write InventoryMonth into Config as text instead of a date')
    ap.add_argument('--check-only', action='store_true', help='click Check setup instead of Build report (no report is written)')
    ap.add_argument('--set', action='append', default=[], metavar='KEY=VALUE',
                    help='type VALUE into the Config setting KEY after the options above (repeatable), e.g. --set ShowSSL=maybe')
    ap.add_argument('--v1-config', action='store_true', help='start from a version 1.3 Config sheet: upgrade, then Check setup')
    ap.add_argument('--v20-config', action='store_true', help='start from a version 2.0 Config sheet: upgrade, then Check setup')
    ap.add_argument('--bas', default=os.path.join(HERE, 'build-open-orders-report.bas'))
    opts = ap.parse_args()
    try:
        import uno
        from com.sun.star.beans import PropertyValue
    except ImportError:
        sys.exit('Python UNO bindings not found (install LibreOffice with python3-uno)')
    if not shutil.which('soffice'):
        sys.exit('soffice (LibreOffice) not found')
    in_folder = os.path.abspath(opts.input_folder)
    out_path = os.path.abspath(opts.output_xlsx)
    size_file = os.path.abspath(opts.size_class_file) if opts.size_class_file else ''
    work = tempfile.mkdtemp(prefix='lo-macro-')
    raw_name = 'macro-output.xlsx'
    msg_log = out_path + '.msg.log'
    if os.path.exists(msg_log):
        os.remove(msg_log)
    code, harness = test_copy(opts.bas, msg_log)
    old = '1.3' if opts.v1_config else ('2.0' if opts.v20_config else '')

    port = free_port()
    proc = subprocess.Popen(['soffice', '--headless', '--invisible', '--norestore', '--nologo', '--nodefault',
                             '-env:UserInstallation=file://' + os.path.join(work, 'profile'),
                             f'--accept=socket,host=127.0.0.1,port={port};urp;'],
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
    def pv(n, v):
        p = PropertyValue(); p.Name = n; p.Value = v; return p
    try:
        ctx = uno.getComponentContext()
        resolver = ctx.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver', ctx)
        for _ in range(90):
            try:
                rctx = resolver.resolve(f'uno:socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext'); break
            except Exception:
                time.sleep(1)
        else:
            sys.exit('LibreOffice did not start')
        desktop = rctx.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop', rctx)

        def uno_call(what, src, dst):
            if what == 'convert':
                doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(src), '_blank', 0, (pv('Hidden', True),))
                doc.storeToURL(uno.systemPathToFileUrl(dst), (pv('FilterName', 'Calc MS Excel 2007 XML'),))
                doc.close(True)

        doc = desktop.loadComponentFromURL('private:factory/scalc', '_blank', 0, (pv('Hidden', False),))
        libs = doc.BasicLibraries
        libs.VBACompatibilityMode = True
        if not libs.hasByName('Standard'):
            libs.createLibrary('Standard')
        lib = libs.getByName('Standard')
        lib.insertByName('modOpenOrdersReport', code)
        lib.insertByName('modLibreOfficeHarness', harness)
        script = doc.getScriptProvider().getScript(
            'vnd.sun.star.script:Standard.modLibreOfficeHarness.RunTest?language=Basic&location=document')
        t0 = time.time()
        script.invoke((in_folder, work, raw_name, size_file, opts.city_column, opts.inventory_month,
                       'TRUE' if opts.show_cancelled else 'FALSE', 'TRUE' if opts.show_ssl else 'FALSE',
                       opts.month_as_text, bool(opts.check_only or old), chr(30).join(opts.set), old), (), ())
        elapsed = time.time() - t0
        lines = open(msg_log).read().split('\n') if os.path.exists(msg_log) else []
        records = [l[len('RECORD|'):] for l in lines if l.startswith('RECORD|')]
        log = '\n'.join(l for l in lines if not l.startswith('RECORD|'))
        print(log.strip() or '(no message boxes)')
        raw = os.path.join(work, raw_name)
        if opts.check_only or old:
            if 'HARNESS ERROR' in log:
                sys.exit('the run failed in the harness (see messages above)')
            print(f'{"upgrade and " if old else ""}setup check ran in LibreOffice in {elapsed:.1f} s; no report built')
            return
        if not os.path.exists(raw) or 'Failed while' in log or 'HARNESS ERROR' in log:
            sys.exit(f'macro run failed after {elapsed:.1f} s (see messages above)')
        for c in list(desktop.Components):
            try:
                c.close(True)
            except Exception:
                pass
        normalise(raw, out_path, uno_call, records)
        print(f'macro ran in LibreOffice in {elapsed:.1f} s; normalised report saved to {out_path}')
    finally:
        try:
            desktop.terminate()
        except Exception:
            pass
        try:
            proc.wait(timeout=20)
        except Exception:
            proc.kill()
        shutil.rmtree(work, ignore_errors=True)

if __name__ == '__main__':
    main()
